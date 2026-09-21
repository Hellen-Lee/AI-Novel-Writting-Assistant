"""附件解析器注册表（D1）与最简解析器实现。

契约见 docs/changes/260913 待办清单 C/D 节：

- D1 注册表：``{kind → extract(file_path) -> ParsedContent}``，新格式即新注册项，
  调用方无感；解析异常一律捕获转 ParseError（中文原因），绝不抛出阻断整批上传。
- D2 纯文本：txt / md，UTF-8 读取、解码失败回退 GBK（中文小说场景常见）。
- D3/D4：docx 提取段落与表格；PDF 提取文本层并识别扫描版/加密文件。
- D6 截断：单文件上限 MAX_TEXT_CHARS，MVP 只保留头部 + 末尾标注。
- C1 图片：Pillow 试开（二道防线，H1）→ 等比缩略图（最长边 256px 转 JPEG），
  EXIF 方向仅纠正缩略图、原图不动；GIF 取第一帧。

统一产物（文件）：``{kind, text, text_truncated, chars}``；
图片产物（不含 thumb_rel，由调用方按存储约定拼接）：``{width, height}``。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Callable, Optional

from PIL import Image, ImageOps
from docx import Document
from pypdf import PdfReader

# --- D6 文本截断策略 ---
MAX_TEXT_CHARS = 8000
TRUNCATED_MARK = "\n…（已截断，共 {chars} 字符）"

# --- C1 缩略图 ---
THUMB_MAX_EDGE = 256
THUMB_JPEG_QUALITY = 85


class ParseError(ValueError):
    """解析失败；message 为可直接展示的中文原因。"""


# ---------------------------------------------------------------------------
# D1 文件解析器注册表
# ---------------------------------------------------------------------------

# kind（扩展名去点）→ extract(file_path) -> {kind, text, text_truncated, chars}
_FILE_PARSERS: dict[str, Callable[[Path], dict[str, Any]]] = {}


def register_file_parser(*kinds: str):
    """注册文件解析器：@register_file_parser("txt", "md")。"""

    def _decorator(fn: Callable[[Path], dict[str, Any]]):
        for kind in kinds:
            _FILE_PARSERS[kind] = fn
        return fn

    return _decorator


def parse_file_attachment(ext: str, path: Path) -> Optional[dict[str, Any]]:
    """按扩展名分发到已注册解析器；未注册的格式返回 None（解析待接入）。

    D5（csv/json）尚未接入，保持 ready + parsed=None。
    """
    kind = ext.lstrip(".").lower()
    parser = _FILE_PARSERS.get(kind)
    if parser is None:
        return None
    return parser(path)


# ---------------------------------------------------------------------------
# D2 纯文本类：txt / md（保留原文，不做 markdown 剥离）
# ---------------------------------------------------------------------------


def _truncate_text(kind: str, text: str) -> dict[str, Any]:
    """D6：头部保留 + 末尾标注；text 为截断后的最终文本。"""
    chars = len(text)
    if chars <= MAX_TEXT_CHARS:
        return {"kind": kind, "text": text, "text_truncated": False, "chars": chars}
    return {
        "kind": kind,
        "text": text[:MAX_TEXT_CHARS] + TRUNCATED_MARK.format(chars=chars),
        "text_truncated": True,
        "chars": chars,
    }


@register_file_parser("txt", "md")
def _extract_plain_text(path: Path) -> dict[str, Any]:
    kind = path.suffix.lstrip(".").lower()
    raw = path.read_bytes()
    text: Optional[str] = None
    for encoding in ("utf-8", "gbk"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ParseError("文本编码无法识别（仅支持 UTF-8 / GBK）")
    return _truncate_text(kind, text)


@register_file_parser("docx")
def _extract_docx(path: Path) -> dict[str, Any]:
    """D3：提取段落与表格，保留文档的基本阅读节奏。"""
    try:
        document = Document(path)
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text
        ]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return _truncate_text("docx", "\n\n".join(parts))
    except Exception as exc:
        raise ParseError("DOCX 文件解析失败，请检查文件是否损坏") from exc


@register_file_parser("pdf")
def _extract_pdf(path: Path) -> dict[str, Any]:
    """D4：提取 PDF 文本层；扫描版 PDF 保持 ready 并给出提示。"""
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
            except Exception as exc:
                raise ParseError("PDF 已加密，请先解除密码") from exc
            if not decrypted:
                raise ParseError("PDF 已加密，请先解除密码")

        page_texts = [(page.extract_text() or "") for page in reader.pages]
        text = "\n\n".join(page_texts).strip()
        if not text:
            return {
                "kind": "pdf",
                "text": "",
                "text_truncated": False,
                "chars": 0,
                "note": "扫描版PDF，无文本层",
            }
        return _truncate_text("pdf", text)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("PDF 文件解析失败，请检查文件是否损坏") from exc


# ---------------------------------------------------------------------------
# C1 图片：试开校验 + 等比缩略图
# ---------------------------------------------------------------------------


def parse_image_attachment(path: Path, thumb_path: Path) -> dict[str, Any]:
    """解析图片 → {width, height}，并生成 JPEG 缩略图到 thumb_path。

    损坏图片（Pillow open/decode 抛错）→ ParseError「图片文件损坏或格式异常」。
    """
    try:
        with Image.open(path) as img:
            if getattr(img, "is_animated", False):
                img.seek(0)  # GIF 取第一帧
            img.load()  # 强制解码，损坏数据在此抛错
            width, height = img.size
            thumb = ImageOps.exif_transpose(img)  # EXIF 方向仅纠正缩略图
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("图片文件损坏或格式异常") from exc

    try:
        if thumb.mode in ("RGBA", "LA", "PA"):
            # 透明通道垫白底，避免转 JPEG 后变黑
            background = Image.new("RGB", thumb.size, (255, 255, 255))
            background.paste(thumb, mask=thumb.getchannel("A"))
            thumb = background
        elif thumb.mode != "RGB":
            thumb = thumb.convert("RGB")
        thumb.thumbnail((THUMB_MAX_EDGE, THUMB_MAX_EDGE))
        thumb_path.parent.mkdir(parents=True, exist_ok=True)
        thumb.save(thumb_path, "JPEG", quality=THUMB_JPEG_QUALITY)
    except Exception as exc:
        thumb_path.unlink(missing_ok=True)
        raise ParseError("图片文件损坏或格式异常") from exc

    return {"width": width, "height": height}
