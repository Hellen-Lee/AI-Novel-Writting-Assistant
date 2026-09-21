"""E1~E4 附件发送链路的离线单元检查。

运行：python backend/tests/test_attachment_send.py
"""

from __future__ import annotations

import unittest
import sys
from tempfile import TemporaryDirectory
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.models.schemas import GenerateRequest
from app.services import (
    attachment_parser,
    attachment_store,
    project_store,
    prompt_builder,
    session_store,
)


class AttachmentSendTest(unittest.TestCase):
    def test_generate_request_requires_text_and_limits_attachments(self) -> None:
        request = GenerateRequest(
            session_id="sess_test",
            user_input="  请总结附件  ",
            attachments=["att_1", "att_2"],
        )
        self.assertEqual(request.user_input, "请总结附件")
        with self.assertRaises(ValueError):
            GenerateRequest(session_id="sess_test", user_input="   ")
        with self.assertRaises(ValueError):
            GenerateRequest(
                session_id="sess_test",
                user_input="请求",
                attachments=["a1", "a2", "a3", "a4", "a5", "a6"],
            )

    def test_materials_keep_order_and_show_fallbacks(self) -> None:
        records = [
            {
                "id": "att_text",
                "type": "file",
                "name": "设定.txt",
                "parsed": {"text": "人物住在旧城区。"},
            },
            {
                "id": "att_image",
                "type": "image",
                "name": "地图.png",
                "parsed": {"width": 100, "height": 80},
            },
        ]
        materials = prompt_builder.format_attachment_materials(records)
        self.assertLess(materials.index("设定.txt"), materials.index("地图.png"))
        self.assertIn("人物住在旧城区", materials)
        self.assertIn("未启用识图能力", materials)

    def test_message_stores_attachment_summary_and_selection(self) -> None:
        message = session_store.make_message(
            "user",
            "请阅读",
            attachments=[{"id": "att_1", "type": "file", "name": "a.txt"}],
            selection="选中的文字",
        )
        self.assertEqual(message["attachments"][0]["id"], "att_1")
        self.assertEqual(message["selection"], "选中的文字")

    def test_attachment_message_survives_session_roundtrip_without_file_text(self) -> None:
        with TemporaryDirectory() as temp_dir:
            projects_dir = Path(temp_dir) / "projects"
            with patch.object(project_store, "PROJECTS_DIR", projects_dir):
                project = project_store.create_project(name="附件回显测试")
                session = session_store.create_session(project["meta"]["id"])
                attachment = {
                    "id": "att_roundtrip",
                    "type": "file",
                    "name": "设定.txt",
                    "mime": "text/plain",
                    "size": 128,
                    "status": "ready",
                    "parsed": {
                        "kind": "txt",
                        "text_truncated": True,
                        "chars": 9000,
                    },
                }
                message = session_store.make_message(
                    "user",
                    "请阅读附件",
                    attachments=[attachment],
                    selection="旧城区",
                )
                session_store.append_messages(
                    project["meta"]["id"],
                    session["id"],
                    [message],
                )

                loaded = session_store.get_session(project["meta"]["id"], session["id"])
                loaded_message = loaded["messages"][0]
                self.assertEqual(loaded_message["attachments"], [attachment])
                self.assertEqual(loaded_message["selection"], "旧城区")
                self.assertNotIn("text", loaded_message["attachments"][0]["parsed"])

                raw = session_store._session_path(
                    project["meta"]["id"], session["id"]
                ).read_text(encoding="utf-8")
                self.assertNotIn("人物住在旧城区", raw)

    def test_docx_parser_keeps_paragraphs_and_tables(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "设定.docx"
            document = Document()
            document.add_paragraph("第一段")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "角色"
            table.cell(0, 1).text = "林舟"
            document.save(path)

            parsed = attachment_parser.parse_file_attachment("docx", path)

            self.assertIsNotNone(parsed)
            assert parsed is not None
            self.assertIn("第一段", parsed["text"])
            self.assertIn("角色\t林舟", parsed["text"])
            self.assertEqual(parsed["chars"], len("第一段\n\n角色\t林舟"))

    def test_pdf_parser_marks_scanned_pdf_and_rejects_encrypted_pdf(self) -> None:
        with TemporaryDirectory() as temp_dir:
            blank_path = Path(temp_dir) / "扫描.pdf"
            PdfWriter().write(blank_path)
            parsed = attachment_parser.parse_file_attachment("pdf", blank_path)
            self.assertIsNotNone(parsed)
            assert parsed is not None
            self.assertEqual(parsed["text"], "")
            self.assertEqual(parsed["note"], "扫描版PDF，无文本层")

            encrypted_path = Path(temp_dir) / "加密.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.encrypt("secret")
            writer.write(encrypted_path)
            with self.assertRaises(attachment_parser.ParseError) as error:
                attachment_parser.parse_file_attachment("pdf", encrypted_path)
            self.assertIn("PDF 已加密", str(error.exception))

    def test_delete_session_removes_attachment_directory_before_session(self) -> None:
        with TemporaryDirectory() as temp_dir:
            projects_dir = Path(temp_dir) / "projects"
            with patch.object(project_store, "PROJECTS_DIR", projects_dir):
                project = project_store.create_project(name="级联删除测试")
                session = session_store.create_session(project["meta"]["id"])
                attachment_dir = attachment_store.attachments_dir(
                    project["meta"]["id"], session["id"]
                )
                attachment_dir.mkdir(parents=True)
                (attachment_dir / "att_test.txt").write_text("内容", encoding="utf-8")

                session_store.delete_session(project["meta"]["id"], session["id"])

                self.assertFalse(attachment_dir.exists())
                self.assertFalse(
                    session_store._session_path(
                        project["meta"]["id"], session["id"]
                    ).exists()
                )
                self.assertEqual(session_store.list_sessions(project["meta"]["id"]), [])


if __name__ == "__main__":
    unittest.main()
